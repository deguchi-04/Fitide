"""Generate the compact PortFIR food index used by Fitide.

Usage: python scripts/extract_portfir.py path/to/insa_tca.xlsx
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl


SOURCE_URL = (
    "https://portfir.insa.min-saude.pt/wp-content/uploads/2025/11/insa_tca.xlsx"
)


def number(value: object) -> float:
    return float(value) if isinstance(value, (int, float)) else 0.0


def main() -> None:
    source = Path(sys.argv[1])
    sheet = openpyxl.load_workbook(source, data_only=True).worksheets[0]
    foods = []
    for row in range(3, sheet.max_row + 1):
        foods.append(
            {
                "id": str(sheet.cell(row, 1).value),
                "name": str(sheet.cell(row, 2).value).strip(),
                "calories": number(sheet.cell(row, 6).value),
                "protein": number(sheet.cell(row, 20).value),
                "carbs": number(sheet.cell(row, 14).value),
                "fat": number(sheet.cell(row, 8).value),
                "fiber": number(sheet.cell(row, 19).value),
                "calcium": number(sheet.cell(row, 47).value),
                "iron": number(sheet.cell(row, 50).value),
                "vitaminC": number(sheet.cell(row, 42).value),
            }
        )

    target = Path(__file__).resolve().parents[1] / "lib" / "portfir.ts"
    payload = json.dumps(foods, ensure_ascii=False, separators=(",", ":"))
    target.write_text(
        "import type { Nutrients } from './fit-types';\n\n"
        "export interface PortfirFood extends Nutrients { id: string; name: string; }\n\n"
        "// Valores por 100 g de parte edível (bebidas alcoólicas: por 100 ml).\n"
        "// Fonte oficial portuguesa: PortFIR/INSA, TCA v7.1 (2026).\n"
        f"export const PORTFIR_URL = '{SOURCE_URL}';\n"
        f"export const portfirFoods: PortfirFood[] = {payload};\n",
        encoding="utf-8",
    )
    print(f"Wrote {len(foods)} foods to {target}")


if __name__ == "__main__":
    main()
